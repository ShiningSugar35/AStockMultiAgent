from __future__ import annotations

import json
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document as OpenXmlDocument

from astock.core.object_store import ObjectStore
from astock.documents.block_repository import DocumentBlockRepository
from astock.documents.media import DOC_MIME, DOCX_MIME, TEXT_MIME
from astock.documents.reflow_parser import OfficialReflowableParseService
from astock.documents.repository import DocumentRepository
from astock.schemas import DocumentBlockKind, DocumentType, SourceDocument, SourceSnapshot


def _docx_bytes() -> bytes:
    document = OpenXmlDocument()
    document.add_heading("1 Management Discussion", level=1)
    document.add_paragraph("Revenue increased from the prior period.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _register(
    state,
    objects: ObjectStore,
    payload: bytes,
    *,
    document_id: str,
    mime: str,
) -> tuple[SourceDocument, SourceSnapshot]:
    object_ref = objects.put_bytes(payload)
    now = datetime(2026, 9, 27, 6, 0, tzinfo=UTC)
    snapshot = SourceSnapshot(
        snapshot_id=f"fixture:{document_id}:{object_ref.sha256}",
        source_id="fixture-official",
        object_sha256=object_ref.sha256,
        fetched_at=now,
        available_to_system_at=now,
        source_url=f"https://example.invalid/{document_id}",
        mime=mime,
        byte_size=object_ref.byte_size,
        rights_status="TEST_FIXTURE",
    )
    state.register_snapshot(snapshot)
    canonical = state.get_snapshot(snapshot.snapshot_id)
    assert canonical is not None
    document = SourceDocument(
        document_id=document_id,
        title="Fixture official document",
        publisher="fixture-official",
        document_type=DocumentType.ANNOUNCEMENT,
        company_ids=["600519"],
        published_at=now,
        effective_at=now,
        disclosure_id=document_id,
        source_url=f"https://example.invalid/{document_id}",
        rights_status="TEST_FIXTURE",
    )
    DocumentRepository(state).register(document, canonical)
    return document, canonical


def _metadata(objects: ObjectStore, block) -> dict[str, object]:
    return json.loads(objects.get_bytes(block.metadata_object_sha256).decode("utf-8"))


def test_docx_parse_persists_stable_heading_and_table_row_locators(tmp_path: Path, state) -> None:
    objects = ObjectStore(tmp_path / "objects")
    document, snapshot = _register(
        state,
        objects,
        _docx_bytes(),
        document_id="fixture:docx",
        mime=DOCX_MIME,
    )
    blocks = DocumentBlockRepository(state)
    service = OfficialReflowableParseService(objects, state, blocks)

    first = service.parse(document, snapshot)
    repeated = service.parse(document, snapshot)

    assert first.cache_hit is False
    assert repeated.cache_hit is True
    assert repeated.block_ids == first.block_ids
    stored = blocks.blocks_for(snapshot.snapshot_id, first.parser_version)
    assert [item.block_kind for item in stored] == [
        DocumentBlockKind.PARAGRAPH,
        DocumentBlockKind.PARAGRAPH,
        DocumentBlockKind.TABLE_ROW,
    ]
    assert stored[0].is_heading is True
    heading_meta = _metadata(objects, stored[0])
    assert heading_meta["section_path"] == ["1 Management Discussion"]
    row_meta = _metadata(objects, stored[2])
    assert row_meta["cell_spans"] == [
        {"cell_index": 1, "char_start": 0, "char_end": 6},
        {"cell_index": 2, "char_start": 7, "char_end": 12},
    ]
    assert objects.get_bytes(stored[2].text_object_sha256).decode("utf-8") == "Metric\tValue"


def test_direct_reflow_cache_repairs_missing_derived_block_object(tmp_path: Path, state) -> None:
    objects = ObjectStore(tmp_path / "objects")
    document, snapshot = _register(
        state,
        objects,
        _docx_bytes(),
        document_id="fixture:docx-cache-repair",
        mime=DOCX_MIME,
    )
    blocks = DocumentBlockRepository(state)
    service = OfficialReflowableParseService(objects, state, blocks)
    first = service.parse(document, snapshot)
    stored = blocks.blocks_for(snapshot.snapshot_id, first.parser_version)
    assert stored
    missing_hash = stored[0].text_object_sha256
    objects.path_for(missing_hash).unlink()
    assert not objects.verify(missing_hash)

    repaired = service.parse(document, snapshot)

    assert repaired.block_ids == first.block_ids
    assert repaired.cache_hit is False
    assert objects.verify(missing_hash)


def test_text_parse_records_line_and_character_offsets(tmp_path: Path, state) -> None:
    objects = ObjectStore(tmp_path / "objects")
    payload = b"alpha\n\nbeta gamma\n"
    document, snapshot = _register(
        state,
        objects,
        payload,
        document_id="fixture:text",
        mime=TEXT_MIME,
    )

    report = OfficialReflowableParseService(objects, state).parse(document, snapshot)
    stored = DocumentBlockRepository(state).blocks_for(snapshot.snapshot_id, report.parser_version)

    assert len(stored) == 2
    first_meta = _metadata(objects, stored[0])
    second_meta = _metadata(objects, stored[1])
    assert (
        first_meta["line_start"],
        first_meta["document_char_start"],
        first_meta["document_char_end"],
    ) == (
        1,
        0,
        5,
    )
    assert (
        second_meta["line_start"],
        second_meta["document_char_start"],
        second_meta["document_char_end"],
    ) == (3, 7, 17)
    assert first_meta["parent_object_sha256"] == snapshot.object_sha256
    assert second_meta["text_encoding"] == "utf-8-sig"


def test_legacy_doc_fixture_is_actually_parsed_with_native_line_offsets(
    tmp_path: Path,
    state,
) -> None:
    fixture = (
        Path(__file__).resolve().parents[2] / "tests/fixtures/documents/legacy_word97_sample.doc"
    )
    payload = fixture.read_bytes()
    assert payload.startswith(bytes.fromhex("D0CF11E0A1B11AE1"))
    objects = ObjectStore(tmp_path / "objects")
    document, snapshot = _register(
        state,
        objects,
        payload,
        document_id="fixture:legacy-doc",
        mime=DOC_MIME,
    )

    report = OfficialReflowableParseService(objects, state).parse(document, snapshot)
    stored = DocumentBlockRepository(state).blocks_for(snapshot.snapshot_id, report.parser_version)

    assert report.parser_name == "legacy-doc"
    assert "legacy-doc-0.2.1" in report.parser_version
    assert len(stored) == 3
    assert [
        objects.get_bytes(block.text_object_sha256).decode("utf-8") for block in stored
    ] == [
        "Legacy DOC actual parser fixture",
        "Revenue 123.45",
        "End",
    ]
    metadata = [_metadata(objects, block) for block in stored]
    assert [item["line_start"] for item in metadata] == [1, 2, 3]
    assert metadata[0]["document_char_start"] == 0
    second_start = metadata[1]["document_char_start"]
    first_end = metadata[0]["document_char_end"]
    assert isinstance(second_start, int) and isinstance(first_end, int)
    assert second_start > first_end
    assert all(item["parent_object_sha256"] == snapshot.object_sha256 for item in metadata)


def test_missing_legacy_doc_parser_does_not_break_docx_or_text_imports(
    tmp_path: Path,
    state,
    monkeypatch,
) -> None:
    import builtins

    original_import = builtins.__import__

    def guarded_import(name, *args, **kwargs):
        if name == "legacy_doc":
            raise ImportError("legacy DOC parser intentionally unavailable")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", guarded_import)
    objects = ObjectStore(tmp_path / "objects")
    service = OfficialReflowableParseService(objects, state)

    docx_document, docx_snapshot = _register(
        state,
        objects,
        _docx_bytes(),
        document_id="fixture:docx-without-legacy-parser",
        mime=DOCX_MIME,
    )
    text_document, text_snapshot = _register(
        state,
        objects,
        b"alpha\nbeta\n",
        document_id="fixture:text-without-legacy-parser",
        mime=TEXT_MIME,
    )
    assert service.parse(docx_document, docx_snapshot).block_ids
    assert service.parse(text_document, text_snapshot).block_ids

    legacy_fixture = (
        Path(__file__).resolve().parents[2] / "tests/fixtures/documents/legacy_word97_sample.doc"
    )
    legacy_document, legacy_snapshot = _register(
        state,
        objects,
        legacy_fixture.read_bytes(),
        document_id="fixture:legacy-doc-without-parser",
        mime=DOC_MIME,
    )
    with pytest.raises(ValueError, match="Legacy DOC parser is unavailable"):
        service.parse(legacy_document, legacy_snapshot)


def test_docx_financial_certification_uses_native_block_evidence(tmp_path: Path, state) -> None:
    from datetime import date
    from decimal import Decimal

    from astock.documents.page_repository import DocumentPageRepository
    from astock.evidence.repository import EvidenceRepository
    from astock.financial_sources.certification import FinancialPdfCertifier
    from astock.financial_sources.config import FinancialFieldMapping
    from astock.financial_sources.official import OfficialFinancialReport
    from astock.schemas import (
        EvidenceLocatorType,
        FinancialDurationSemantics,
        FinancialFieldCode,
        FinancialPeriodType,
        FinancialSourceObservation,
        FinancialStatementScope,
        FinancialStatementType,
        FinancialUnit,
        InstrumentType,
        Market,
        OfficialFinancialLineageKind,
    )

    source = OpenXmlDocument()
    source.add_paragraph("合并资产负债表")
    source.add_paragraph("2025年12月31日")
    source.add_paragraph("币种：人民币")
    source.add_paragraph("单位：万元")
    table = source.add_table(rows=2, cols=3)
    for index, text in enumerate(("项目", "2025年12月31日", "2024年12月31日")):
        table.cell(0, index).text = text
    for index, text in enumerate(("资产总计", "123.00", "100.00")):
        table.cell(1, index).text = text
    buffer = BytesIO()
    source.save(buffer)

    objects = ObjectStore(tmp_path / "objects")
    document, snapshot = _register(
        state,
        objects,
        buffer.getvalue(),
        document_id="fixture:financial-docx",
        mime=DOCX_MIME,
    )
    parser = OfficialReflowableParseService(objects, state)
    parsed = parser.parse(document, snapshot)
    pages_repo = DocumentPageRepository(state)
    blocks_repo = DocumentBlockRepository(state)
    pages = [pages_repo.get_page_by_id(page_id) for page_id in parsed.page_ids]
    blocks = [blocks_repo.get_block_by_id(block_id) for block_id in parsed.block_ids]
    assert all(page is not None for page in pages)
    assert all(block is not None for block in blocks)

    report = OfficialFinancialReport(
        document=document,
        index_snapshot=snapshot,
        lineage_kind=OfficialFinancialLineageKind.LEGACY_UNVERIFIED,
        lineage_snapshot_ids=[],
        exhaustive_proof_allowed=False,
        snapshot=snapshot,
        pages=[page for page in pages if page is not None],
        blocks=[block for block in blocks if block is not None],
        supersedes_document_id=None,
    )
    mapping = FinancialFieldMapping(
        field_code=FinancialFieldCode.TOTAL_ASSETS,
        statement_type=FinancialStatementType.BALANCE_SHEET,
        official_label="资产总计",
        provider_fields={"fixture-secondary": "total_assets"},
        unit=FinancialUnit.TEN_THOUSAND_CNY,
    )
    observation = FinancialSourceObservation(
        observation_id="1" * 64,
        company_id="600519",
        instrument_id="XSHG:600519",
        market=Market.XSHG,
        instrument_type=InstrumentType.STOCK,
        instrument_release_id="2" * 64,
        instrument_manifest_artifact_id="market-reference:" + "3" * 64,
        instrument_manifest_object_hash="4" * 64,
        instrument_content_hash="5" * 64,
        period_start=None,
        period_end=date(2025, 12, 31),
        period_type=FinancialPeriodType.ANNUAL,
        duration_semantics=FinancialDurationSemantics.INSTANT,
        statement_type=FinancialStatementType.BALANCE_SHEET,
        statement_scope=FinancialStatementScope.CONSOLIDATED,
        field_code=FinancialFieldCode.TOTAL_ASSETS,
        provider_field="total_assets",
        reported_value=Decimal("123.00"),
        unit=FinancialUnit.TEN_THOUSAND_CNY,
        provider_id="fixture-secondary",
        source_snapshot_id=snapshot.snapshot_id,
        source_request_hash="6" * 64,
        available_to_system_at=snapshot.available_to_system_at,
    )

    facts, reasons = FinancialPdfCertifier(state, objects).certify(
        report,
        [observation],
        [mapping],
    )

    assert reasons == []
    assert len(facts) == 1
    evidence = EvidenceRepository(state).get_evidence(facts[0].evidence_ids[0])
    assert evidence is not None
    assert evidence.locator.locator_type is EvidenceLocatorType.BLOCK_TEXT
    assert evidence.page_id is None
    assert evidence.block_id is not None
    block = blocks_repo.get_block_by_id(evidence.block_id)
    assert block is not None
    assert block.block_kind is DocumentBlockKind.TABLE_ROW
    metadata = _metadata(objects, block)
    assert metadata["table_index"] if "table_index" in metadata else block.table_index == 1
    excerpt = objects.get_bytes(evidence.excerpt_object_sha256).decode("utf-8")
    assert "资产总计" in excerpt
    assert "123.00" in excerpt
